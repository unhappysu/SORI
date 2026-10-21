"""process_measurements.py 검산 — 손으로 확인할 수 있는 단순 사례.

※ 본 테스트는 생성형 AI(Claude)를 활용하여 작성되었음.

    python -m pytest field_survey/tests -q
"""
import csv
import datetime as dt
import math
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import build_fieldbook_xlsx  # noqa: E402
import process_measurements as pm  # noqa: E402


# ───────────────────────── 계산식
def test_energy_average_same_levels():
    assert pm.energy_average([70.0, 70.0]) == pytest.approx(70.0)


def test_energy_average_60_70():
    # 10·lg((10^6 + 10^7)/2) = 10·lg(5.5×10^6) = 67.40
    assert pm.energy_average([60.0, 70.0]) == pytest.approx(10 * math.log10(5.5e6))
    assert pm.round_half_up(pm.energy_average([60.0, 70.0]), 1) == 67.4


def test_energy_average_time_weighted():
    # (20분×10^6 + 10분×10^7) / 30분 = 1.2×10^8 / 30 = 4×10^6 → 10·lg(4×10^6) = 66.02 dB
    assert pm.energy_average([60.0, 70.0], [20.0, 10.0]) == pytest.approx(10 * math.log10(4e6))
    assert pm.round_half_up(pm.energy_average([60.0, 70.0], [20.0, 10.0]), 1) == 66.0


def test_energy_not_below_arithmetic():
    levels = [58.0, 63.0, 71.0]
    assert pm.energy_average(levels) >= sum(levels) / len(levels)


@pytest.mark.parametrize("x, nd, expect", [(65.5, 0, 66), (64.5, 0, 65), (68.25, 1, 68.3), (71.04, 1, 71.0)])
def test_round_half_up(x, nd, expect):
    assert pm.round_half_up(x, nd) == expect


def test_arithmetic_average_legal():
    # (65 + 66)/2 = 65.5 → 최종 66 (사사오입; 파이썬 round(65.5) = 66, round(64.5) = 64 와 구분)
    assert pm.arithmetic_average_legal([65.0, 66.0]) == (65.5, 66)
    assert pm.arithmetic_average_legal([64.0, 65.0]) == (64.5, 65)


def test_minutes_between_midnight():
    assert pm.minutes_between(dt.time(23, 55), dt.time(0, 5)) == pytest.approx(10.0)
    assert pm.minutes_between(dt.time(7, 30), dt.time(7, 40, 30)) == pytest.approx(10.5)


def test_diff_stats():
    # 차이 1, −1, 3 → 평균 1, 표본표준편차 √((0² + 2² + 2²)/2) = 2, 최대 |3|
    s = pm.diff_stats([("A", 1.0), ("B", -1.0), ("C", 3.0)])
    assert (s["n"], s["mean_diff"], s["sd_diff"], s["max_abs_diff"], s["max_abs_point"]) == (3, 1.0, 2.0, 3.0, "C")


def test_empty_inputs_raise():
    with pytest.raises(ValueError):
        pm.energy_average([])
    with pytest.raises(ValueError):
        pm.arithmetic_average_legal([])


# ───────────────────────── 야장 → CSV
def _fieldbook(tmp_path):
    src = build_fieldbook_xlsx.build(tmp_path / "야장.xlsx")
    wb = load_workbook(src)
    p = wb["2_측정점"]
    p["F5"], p["G5"] = 205100.0, 553120.0          # A1 좌표(지도 판독값을 덮어씀)
    n = wb["3_소음기록"]
    rows = [("A1", "D1", 70.0, 10, "Y"), ("A1", "D2", 72.0, 10, "Y"), ("A1", "D3", 68.0, 10, "Y"),
            ("A1", "D4", 74.0, 10, "Y"), ("A1", "N1", 62.0, 10, "Y"), ("A1", "N2", 58.0, 10, "Y"),
            ("A2", "D1", 65.0, 10, "Y"), ("A2", "D2", 67.0, 10, "N"), ("A2", "D3", 66.0, 20, "Y")]
    for i, (pid, rd, L, mins, ok) in enumerate(rows):
        r = 5 + i
        start = dt.datetime(2026, 10, 21, 7, 30)
        _put(n, r, 측정점=pid, 회차=rd, LAeq=L, **{"유효(Y/N)": ok, "교통량 참조": f"{rd}_TC",
             "시작 시각": start.time(), "종료 시각": (start + dt.timedelta(minutes=mins)).time()})
    wb.save(src)
    return src


def _put(ws, r, **values):
    """3_소음기록 r행에 열 이름으로 값 넣기(열 위치가 바뀌어도 그대로 씀)"""
    for name, v in values.items():
        ws[f"{build_fieldbook_xlsx.NC[name]}{r}"] = v


def _read(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return {(r["point_id"], r.get("period")): r for r in csv.DictReader(f)}


def test_fieldbook_to_measured_levels(tmp_path):
    out = tmp_path / "data"
    pm.main([str(_fieldbook(tmp_path)), "--out", str(out)])
    m = _read(out / "measured_levels.csv")
    assert list(csv.reader(open(out / "measured_levels.csv", encoding="utf-8-sig")))[0][:9] == \
        ["point_id", "x", "y", "z", "period", "Leq", "n_samples", "valid_flag", "traffic_count_ref"]
    r1d = m[("A1", "day")]
    # 에너지평균 10·lg((10^7 + 10^7.2 + 10^6.8 + 10^7.4)/4) = 71.56 → 71.6, 산술평균 71.0
    assert (float(r1d["Leq"]), int(r1d["n_samples"]), r1d["valid_flag"]) == (71.6, 4, "Y")
    assert (float(r1d["Leq_legal_mean"]), int(r1d["Leq_legal_final"]), r1d["legal_count_ok"]) == (71.0, 71, "Y")
    assert (float(r1d["x"]), float(r1d["y"]), float(r1d["z"])) == (205100.0, 553120.0, 1.5)
    assert r1d["traffic_count_ref"] == "D1_TC;D2_TC;D3_TC;D4_TC"
    r1n = m[("A1", "night")]
    # 10·lg((10^6.2 + 10^5.8)/2) = 60.44 → 60.4, 산술평균 60.0, 밤 2회 충족
    assert (float(r1n["Leq"]), float(r1n["Leq_legal_mean"]), r1n["legal_count_ok"]) == (60.4, 60.0, "Y")
    f1d = m[("A2", "day")]
    # 무효(N) 행 제외. (10분×10^6.5 + 20분×10^6.6)/30분 → 65.69 → 65.7, 산술 (65+66)/2 = 65.5 → 66, 낮 4회 미달
    assert (float(f1d["Leq"]), int(f1d["n_samples"]), float(f1d["Leq_legal_mean"]), int(f1d["Leq_legal_final"]),
            f1d["legal_count_ok"], float(f1d["total_minutes"])) == (65.7, 2, 65.5, 66, "N", 30.0)
    assert ("A2", "night") not in m


def test_short_measurement_flagged(tmp_path):
    src = _fieldbook(tmp_path)
    wb = load_workbook(src)
    n = wb["3_소음기록"]
    _put(n, 5, **{"종료 시각": dt.time(7, 35)})       # A1 D1 을 5분으로
    wb.save(src)
    out = tmp_path / "d2"
    pm.main([str(src), "--out", str(out)])
    assert _read(out / "measured_levels.csv")[("A1", "day")]["valid_flag"] == "N"


def test_model_comparison(tmp_path):
    model = tmp_path / "model.csv"
    model.write_text("point_id,period,Leq_calc\nA1,day,70.0\nA2,day,66.7\nA1,night,\n", encoding="utf-8")
    out = tmp_path / "data"
    pm.main([str(_fieldbook(tmp_path)), "--model", str(model), "--out", str(out)])
    c = _read(out / "comparison.csv")
    assert float(c[("A1", "day")]["diff"]) == 1.6       # 71.6 − 70.0
    assert float(c[("A2", "day")]["diff"]) == -1.0      # 65.7 − 66.7
    assert c[("A1", "night")]["note"] == "계산값 없음"
    s = next(csv.DictReader(open(out / "comparison_summary.csv", encoding="utf-8-sig")))
    # 차이 1.6, −1.0 → 평균 0.3, 표준편차 √((1.3² + 1.3²)/1) = 1.84
    assert (int(s["n"]), float(s["mean_diff"]), float(s["sd_diff"]), float(s["max_abs_diff"]), s["max_abs_point"]) == \
        (2, 0.3, 1.84, 1.6, "A1/day")


def test_model_leq_calc_not_number(tmp_path):
    """Leq_calc 가 숫자가 아니면 행 번호와 값을 담은 오류. 빈칸·nan 은 '계산값 없음'으로 허용"""
    model = tmp_path / "model.csv"
    model.write_text("point_id,period,Leq_calc\nA1,day,70.0\nA2,day,약 66\nA1,night,\nB1,day,nan\n",
                     encoding="utf-8")
    with pytest.raises(ValueError) as e:
        pm.read_model(model)
    assert "3행" in str(e.value) and "'약 66'" in str(e.value) and "A2" in str(e.value)
    model.write_text("point_id,period,Leq_calc,quantity\nA1,day,70.0,LrT\nA1,night,,LrN\n", encoding="utf-8")
    assert pm.read_model(model) == {("A1", "day"): 70.0, ("A1", "night"): None}


def test_extra_columns_do_not_break_reading(tmp_path):
    """3_소음기록 앞쪽에 열을 끼워 넣어도(열 위치 이동) 열 이름으로 읽으므로 결과가 같음"""
    src = _fieldbook(tmp_path)
    pm.main([str(src), "--out", str(tmp_path / "a")])
    wb = load_workbook(src)
    for name in ("3_소음기록", "2_측정점"):
        ws = wb[name]
        ws.insert_cols(2)
        ws.cell(4, 2, "추가 메모 열")
    wb.save(src)
    pm.main([str(src), "--out", str(tmp_path / "b")])
    assert _read(tmp_path / "a" / "measured_levels.csv") == _read(tmp_path / "b" / "measured_levels.csv")


def test_excel_formulas_match_script(tmp_path):
    """엑셀 6_계산 시트 수식 결과 = process_measurements.py 결과. 시각이 빠진 행(유효 Y, LAeq 있음)은 둘 다 제외.
    엑셀 수식 계산에 formulas 패키지가 필요(없으면 건너뜀)."""
    formulas = pytest.importorskip("formulas")
    src = _fieldbook(tmp_path)
    wb = load_workbook(src)
    _put(wb["3_소음기록"], 14, 측정점="A1", 회차="D2", LAeq=90.0, **{"유효(Y/N)": "Y",   # 종료 시각 누락
                                                               "시작 시각": dt.time(11, 30)})
    _put(wb["3_소음기록"], 15, 측정점="A2", 회차="N1", LAeq=55.0, **{"유효(Y/N)": "Y"})  # 시각 둘 다 누락
    wb.save(src)
    pm.main([str(src), "--out", str(tmp_path / "d")])
    m = _read(tmp_path / "d" / "measured_levels.csv")
    sol = formulas.ExcelModel().loads(str(src)).finish().calculate()
    vals = {}
    for k, v in sol.items():
        if "!" in k:
            sh, ref = k.upper().rsplit("!", 1)
            vals[(sh.split("]")[-1].strip("'"), ref)] = v.value[0][0] if hasattr(v, "value") else v
    xl = lambda ref: vals[("6_계산", ref)]  # noqa: E731
    for row, pid in ((5, "A1"), (6, "A2")):          # 6_계산: 측정점 순서 A1, A2
        assert xl(f"A{row}") == pid
        for per, (n_, ar, en) in (("day", "BCF"), ("night", "GHK")):
            got = m.get((pid, per))
            if got is None:
                assert float(xl(f"{n_}{row}")) == 0 and xl(f"{ar}{row}") == "" and xl(f"{en}{row}") == ""
                continue
            assert float(xl(f"{n_}{row}")) == int(got["n_samples"])
            assert float(xl(f"{ar}{row}")) == pytest.approx(float(got["Leq_legal_mean"]))
            assert float(xl(f"{en}{row}")) == pytest.approx(float(got["Leq"]))


def test_same_round_rows_combined(tmp_path):
    """같은 회차에 10분 행 3개를 적으면 회차값 = 30분 에너지평균, 법정 횟수는 회차 수로 셈"""
    src = build_fieldbook_xlsx.build(tmp_path / "야장.xlsx")
    wb = load_workbook(src)
    n = wb["3_소음기록"]
    start = dt.datetime(2026, 10, 21, 7, 30)
    for i, (rd, L, k) in enumerate([("D1", 70.0, 0), ("D1", 73.0, 1), ("D1", 76.0, 2), ("D2", 71.0, 0)]):
        t0 = start + dt.timedelta(minutes=10 * k)
        _put(n, 5 + i, 측정점="A1", 회차=rd, LAeq=L, **{"유효(Y/N)": "Y", "시작 시각": t0.time(),
                                                   "종료 시각": (t0 + dt.timedelta(minutes=10)).time()})
    wb.save(src)
    out = tmp_path / "d"
    pm.main([str(src), "--out", str(out)])
    r = {(x["point_id"], x["round"]): x for x in csv.DictReader(open(out / "measured_rounds.csv", encoding="utf-8-sig"))}
    # D1: 10·lg((10^7 + 10^7.3 + 10^7.6)/3) = 10·lg(6.976×10^7 / 3) = 73.66 → 73.7 dB, 30분
    assert float(r[("A1", "D1")]["Leq"]) == 73.7 and float(r[("A1", "D1")]["minutes"]) == 30.0
    m = _read(out / "measured_levels.csv")[("A1", "day")]
    # 법정: 회차값을 소수 첫째 자리로 한 뒤 (73.7 + 71.0)/2 = 72.35 → 72.4, 회차 2개
    assert int(m["n_samples"]) == 2 and float(m["Leq_legal_mean"]) == 72.4


def test_model_comparison_by_round(tmp_path):
    """모델 CSV 에 round 열이 있으면 회차별로 비교"""
    model = tmp_path / "model.csv"
    model.write_text("point_id,period,round,Leq_calc\nA1,day,D1,69.0\nA1,day,D2,71.0\nA2,day,D1,66.0\n", encoding="utf-8")
    out = tmp_path / "data"
    pm.main([str(_fieldbook(tmp_path)), "--model", str(model), "--out", str(out)])
    c = {(r["point_id"], r["round"]): r for r in csv.DictReader(open(out / "comparison.csv", encoding="utf-8-sig"))}
    assert float(c[("A1", "D1")]["diff"]) == 1.0      # 70.0 − 69.0
    assert float(c[("A1", "D2")]["diff"]) == 1.0      # 72.0 − 71.0
    assert float(c[("A2", "D1")]["diff"]) == -1.0     # 65.0 − 66.0
    assert c[("A1", "N1")]["note"] == "계산값 없음"
