"""인쇄용 소음측정 야장(A4 세로) 생성 — 현장에서 펜으로 쓰는 양식.

※ 본 스크립트는 생성형 AI(Claude)를 활용하여 작성되었음.

    python field_survey/scripts/build_fieldbook_print.py --out 폴더 [--pdf]

내용은 소음측정_야장.xlsx 의 입력 칸과 같다(계산 시트는 인쇄하지 않음).
--pdf: Pages(macOS)로 PDF를 만든다. 맑은 고딕은 macOS 시스템 글꼴이 아니므로 PDF용은
시스템 글꼴 Apple SD Gothic Neo 로 만든다.
"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx_common  # noqa: E402
from docx.shared import Mm  # noqa: E402
from docx_common import AI_NOTE, NOTE_FILL, h, new_document, note_box, page_break, para, table  # noqa: E402
from survey_points import DEVICES, POINTS, TEAMS, VEH  # noqa: E402

ROUNDS = ["D1", "D2", "D3", "D4", "N1", "N2"]
N_CARD_PAGES = 3       # 소음 기록 카드 쪽 수(쪽당 2장) — 필요한 만큼 복사해 씀
LABEL = NOTE_FILL


def merge_row(t, r, c0, c1):
    """r행 c0~c1 칸을 합치고, 합치며 생긴 빈 문단을 지운다."""
    m = t.cell(r, c0).merge(t.cell(r, c1))
    ps = m.paragraphs
    for p in ps[1:]:
        if not p.text.strip():
            p._p.getparent().remove(p._p)
    return m


def head_block(doc, title, fields):
    """쪽 머리: 제목 + 공통 기입란(라벨/값 쌍 3개씩)."""
    h(doc, title, 1)
    rows = []
    for i in range(0, len(fields), 3):
        trio = fields[i:i + 3] + [""] * (3 - len(fields[i:i + 3]))
        rows.append([trio[0], "", trio[1], "", trio[2], ""])
    fills = {(r, c): LABEL for r in range(len(rows)) for c in (0, 2, 4) if rows[r][c]}
    table(doc, rows, [20, 35, 20, 35, 20, 36], header_rows=0, size=9.5, min_height_mm=10, fills=fills)


def page_overview(doc):
    h(doc, "소음측정 야장 — 현장 기록용", 1)
    note_box(doc, f"**{AI_NOTE}.** 숫자는 소수 첫째 자리까지, 시각은 기준 시계(초 단위)로 적는다. "
                  "모르는 칸은 비워 두고 추정값을 적을 때는 '추정'이라고 쓴다. 기입 후 소음측정_야장.xlsx에 옮겨 적는다.")
    h(doc, "1. 측정 개요", 2)
    t1, t2 = list(TEAMS.items())
    rows = [["측정일", "", "요일", ""], ["예비일 사용·이유", "", "기준 시계", ""],
            [f"{t1[0]}({t1[1][0]}) 측정자", "", f"{t2[0]}({t2[1][0]}) 측정자", ""],
            [f"{t1[0]} 교통량 계수", "", f"{t2[0]} 교통량 계수", ""],
            [f"{t1[0]} 기록·안전", "", f"{t2[0]} 기록·안전", ""], ["기상 개요", "", "총괄(회차 연락)", ""]]
    table(doc, rows, [30, 53, 30, 53], header_rows=0, size=9.5, min_height_mm=11,
          fills={(r, c): LABEL for r in range(len(rows)) for c in (0, 2)})
    h(doc, "장비와 시각 동기", 2)
    rows = [["장비 ID", "모델", "일련번호", "등급", "1초 L_{Aeq} 저장", "시각 오차(초)\n기준−장비"]]
    rows += [[d, "", "", "", "☐ 예  ☐ 아니오", ""] for d in DEVICES]
    rows += [["교정기", "", "", "표시값", "dB", ""]]
    table(doc, rows, [22, 34, 32, 18, 34, 26], size=9.5, min_height_mm=11, align="cccccc")
    para(doc, "시각 동기: 측정 시작 전과 하루 끝에 각 장비 화면과 기준 시계를 함께 찍은 사진 번호를 적는다 → 시작 ______  끝 ______",
         size=9.5)
    page_break(doc)


def page_calibration(doc):
    h(doc, "2. 교정 기록", 1)
    para(doc, "회차 시작 전 교정, 회차 끝난 뒤 점검. 단위 dB. 교정 전후 허용 편차는 공정시험기준에 규정이 없어(미확인) "
              "판정 기준은 개요 시트에 따로 정한다.", size=9.5)
    rows = [["회차", "장비", "날짜·시각", "교정 전\n지시값", "교정 후\n지시값", "회차 후\n지시값", "메모"]]
    for rd in ROUNDS:
        for d in DEVICES:
            rows.append([rd, d, "", "", "", "", ""])
    table(doc, rows, [14, 18, 30, 22, 22, 22, 38], size=9, min_height_mm=11, align="ccccccc")
    page_break(doc)


def page_points(doc):
    h(doc, "3. 측정점 정보", 1)
    para(doc, "좌표는 스마트폰 GPS(WGS84) 소수 6자리. 거리는 거리측정기·줄자로 잰 값(m). 계획 위치(지도 판독): "
              + "; ".join(f"{p['point_id']} {p['lat_wgs84']}, {p['lon_wgs84']} — {p['location_desc'].split('.')[0]}" for p in POINTS),
         size=9)
    rows = [["ID", "경도(WGS84)", "위도(WGS84)", "높이\n(m)", "차도 끝\n거리(m)", "도로 중심선\n거리(m)",
             "반사면\n이격(m)", "사진 번호"]]
    for p in POINTS:
        rows.append([p["point_id"], "", "", "", "", "", "", ""])
    table(doc, rows, [12, 30, 30, 14, 18, 20, 18, 24], size=9, min_height_mm=12, align="cccccccc")
    h(doc, "주변 상황(반사면·차폐물·주요 소음원·지역 구분)", 2)
    rows = [["ID", "주변 상황"]] + [[p["point_id"], ""] for p in POINTS]
    table(doc, rows, [12, 154], size=9, min_height_mm=12.5, align="cl")
    page_break(doc)


def noise_card(doc, no):
    rows = [
        ["측정점", "", "회차", "", "장비", ""],
        ["날짜", "", "시작 시각", "", "종료 시각", ""],
        ["측정자", "", "", "", "방풍망", "☐ Y    ☐ N"],
        ["L_{Aeq}", "", "L_{Amax}", "", "L_{Amin}", ""],
        ["L_{10}", "", "L_{50}", "", "L_{90}", ""],
        ["풍속(m/s)", "", "배경소음(참고)", "", "유효", "☐ Y    ☐ N"],
        ["이상 소음\n(시각·코드)", "", "", "", "", ""],
        ["무효 사유·메모", "", "", "", "", ""],
    ]
    fills = {(r, c): LABEL for r in range(6) for c in (0, 2, 4) if rows[r][c]}
    fills.update({(6, 0): LABEL, (7, 0): LABEL})
    t = table(doc, rows, [24, 31, 24, 31, 24, 32], header_rows=0, size=9.5, min_height_mm=10.5, fills=fills)
    merge_row(t, 2, 1, 3)                                  # 측정자 이름 칸
    for r, hmm in ((6, 26), (7, 12)):
        merge_row(t, r, 1, 5)
        t.rows[r].height = Mm(hmm)
    return t


def pages_noise(doc):
    for k in range(N_CARD_PAGES):
        h(doc, f"4. 소음 기록 카드 ({k + 1}/{N_CARD_PAGES} — 필요한 만큼 복사)", 1)
        para(doc, "단위 dB(A). 1회 10분. 이상 소음 코드: H 경적, S 사이렌, M 이륜차 급가속, C 공사, V 대화·확성기, "
                  "A 항공기, B 버스 정차, O 기타.", size=9)
        for i in range(2):
            noise_card(doc, k * 2 + i + 1)
        page_break(doc)


def page_traffic(doc):
    head_block(doc, "5. 교통량 계수표 (5분 단위, 한 회차)", ["회차", "날짜", "팀·측정점", "도로", "방향 ①", "방향 ②"])
    para(doc, "한 장 = 한 회차(30분) × 측정점 앞 도로 양방향. 방향 ①·②에 진행 방향(예: 동행·서행)을 적는다. "
              "이륜차는 따로 센다. 세지 못한 구간은 0이 아니라 빗금(/)으로 표시.", size=9)
    rows = [["구간", "방향", *VEH, "정체\n원활/서행/정체"]]
    for k in range(6):
        for d in ("①", "②"):
            rows.append([f"+{k * 5}~{k * 5 + 5}분" if d == "①" else "", d, "", "", "", "", "", "", ""])
    table(doc, rows, [18, 12, 19, 19, 19, 19, 19, 19, 22], size=8.5, min_height_mm=13, align="c" * 9)
    page_break(doc)


def page_signal_weather(doc):
    h(doc, "6. 신호 주기·주행속도", 1)
    rows = [["회차", "시각", "주기 길이(s)", "녹색(s)\n사가정로 서측", "녹색(s)\n사가정로 동측", "녹색(s)\n전농로 남측",
             "녹색(s)\n전농로25길 북측"]]
    rows += [[""] * 7 for _ in range(8)]
    table(doc, rows, [16, 20, 22, 27, 27, 27, 27], size=9, min_height_mm=9.5, align="ccccccc")
    para(doc, "주행속도(자체 방법): 정한 구간 거리를 신호에 막히지 않은 차량이 지나는 시간. 속도 = 3.6 × 거리 / 시간(km/h).",
         size=9)
    rows = [["회차", "도로·접근로", "방향", "구간 거리(m)", "통과 시간(s)", "통과 시간(s)", "통과 시간(s)", "통과 시간(s)"]]
    rows += [[""] * 8 for _ in range(10)]
    table(doc, rows, [16, 30, 16, 22, 20, 20, 21, 21], size=9, min_height_mm=9.5, align="cccccccc")
    page_break(doc)
    h(doc, "7. 기상 기록", 1)
    para(doc, "풍속은 마이크로폰 높이. 2 m/s 이상 방풍망 필수, 5 m/s 초과 측정 금지(공정시험기준). "
              "비·노면 젖음 시 연기(본 조사 자체 기준).", size=9)
    rows = [["날짜", "시각", "위치·측정점", "풍속\n(m/s)", "풍향", "기온\n(°C)", "습도\n(%)", "노면\n건조/젖음", "강우\n없음/있음"]]
    rows += [[""] * 9 for _ in range(20)]
    table(doc, rows, [18, 16, 26, 16, 16, 16, 16, 21, 21], size=9, min_height_mm=10, align="ccccccccc")


def build(out_path):
    doc = new_document("소음측정 야장(인쇄용)")
    page_overview(doc)
    page_calibration(doc)
    page_points(doc)
    pages_noise(doc)
    page_traffic(doc)
    page_signal_weather(doc)
    docx_common.trim_trailing(doc)
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    out_dir = Path(sys.argv[sys.argv.index("--out") + 1]) if "--out" in sys.argv else Path(__file__).resolve().parents[1] / "forms"
    out_dir.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix="fs_pdf_"))             # 임시 docx는 저장소 밖에
    if "--pdf" in sys.argv:
        docx_common.set_font("Apple SD Gothic Neo")
    src = build(tmp / "fieldbook_print.docx")
    print("docx:", src)
    if "--pdf" in sys.argv:
        pdf = out_dir / "소음측정_야장_인쇄용.pdf"
        docx_common.export_pdf(src, pdf, tmp)
        print("PDF:", pdf)
