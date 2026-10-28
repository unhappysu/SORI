"""측정 야장(xlsx)을 읽어 측정점별 L_Aeq 를 계산하고, 모델 계산값과 비교한다.

※ 본 스크립트는 생성형 AI(Claude)를 활용하여 작성되었음.

사용법
    python field_survey/scripts/process_measurements.py 야장.xlsx
    python field_survey/scripts/process_measurements.py 야장.xlsx --model 모델결과.csv --out field_survey/data

입력
    야장.xlsx   : build_fieldbook_xlsx.py 로 만든 소음측정_야장.xlsx (엑셀 수식 결과는 쓰지 않고 입력 칸만 읽어 다시 계산)
    모델결과.csv : (선택) point_id, period(day/night), Leq_calc 열. SORI 의 point_levels_for_comparison.csv 를
                  그대로 쓸 수 있다. round 열(D1…N2)을 더하면 그 회차의 동시 교통량으로 계산한 값과 회차별로 비교한다

출력(--out 폴더, 기본 field_survey/data)
    measured_rounds.csv  : 측정점·회차별 값(같은 회차의 여러 행은 측정시간 가중 에너지평균으로 합침)
    measured_levels.csv  : point_id, x, y, z, period, Leq, n_samples(회차 수), valid_flag, traffic_count_ref
                           + Leq_legal_mean, Leq_legal_final, legal_count_ok, total_minutes
    comparison.csv       : (--model 이 있을 때) point_id, period, Leq_meas, Leq_calc, diff, note
    comparison_summary.csv : n, mean_diff, sd_diff, max_abs_diff, max_abs_point

계산식과 출처
    1) 측정시간 가중 에너지평균(시뮬레이션 비교용, 본 조사 자체 설정)
           L = 10·lg( Σ T_i·10^(L_i/10) / Σ T_i )
       등가소음도 정의(ES 03300.c 2.10: 총 에너지를 같은 시간의 정상소음 에너지로 등가)에서 나온 식이다.
       여러 회차를 이렇게 합치는 것은 공정시험기준의 법정 방법이 아니다. → 열 Leq
    2) 산술평균(법정 평가용) L = (L_1 + … + L_n) / n
       ES 03301.1c 5.3(환경기준), ES 03304.1d 7.1.1.2(도로교통소음관리기준).
       계산 과정은 소수 첫째 자리, 최종값은 소수 첫째 자리에서 반올림(두 방법 7.1). → 열 Leq_legal_mean, Leq_legal_final
       반올림은 사사오입(0.5는 올림)으로 한다. 파이썬 round()의 짝수 반올림을 쓰지 않는다.
    3) 횟수 요건: 낮 4회 이상(ES 03301.1c 5.3.1), 밤 2회 이상(5.3.2) → legal_count_ok
       회차 간 시간 간격(2시간·4시간) 요건은 이 스크립트에서 확인하지 않는다.

배경소음 보정
    TODO(구현하지 않음): 적용 방법(ES 03301.1c, ES 03304.1d)에 배경소음 보정 조항이 없다.
    보정표는 생활소음 방법(ES 03303.1d 7.2)에만 있어 이 조사에는 쓰지 않는다(2026-10-19 원문 확인).

유효 행
    3_소음기록 에서 유효(Y/N) = "Y" 이고 LAeq 가 숫자이며 측정시간(종료−시작)이 0보다 큰 행만 쓴다.
    valid_flag = "Y": 모든 회차의 측정시간(같은 회차 행의 합)이 10분 이상(ES 03304.1d 7.1.1.1) / 그 밖은 "N".
    야장에는 회차마다 1행(30분 등가소음도)을 적는 것을 원칙으로 한다. 여러 행이면 위와 같이 합친다.
"""
import argparse
import csv
import datetime as dt
import math
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import load_workbook

HEAD_ROW = 4
PERIOD = {"낮": "day", "밤": "night"}
MIN_MINUTES = 10.0
LEGAL_COUNT = {"day": 4, "night": 2}


# ───────────────────────── 계산식
def round_half_up(x, ndigits=0):
    """사사오입. round_half_up(65.5) = 66, round_half_up(64.45, 1) = 64.5"""
    q = Decimal(1).scaleb(-ndigits)
    return float(Decimal(repr(x)).quantize(q, rounding=ROUND_HALF_UP))


def energy_average(levels, durations=None):
    """측정시간 가중 에너지평균 L = 10·lg(Σ T_i·10^(L_i/10) / Σ T_i). durations 가 없으면 같은 시간."""
    if not levels:
        raise ValueError("levels 가 비어 있음")
    if durations is None:
        durations = [1.0] * len(levels)
    if len(durations) != len(levels) or any(t <= 0 for t in durations):
        raise ValueError("durations 는 levels 와 길이가 같고 모두 양수여야 함")
    e = sum(t * 10 ** (L / 10) for L, t in zip(levels, durations))
    return 10 * math.log10(e / sum(durations))


def arithmetic_average_legal(levels):
    """법정 산술평균. (소수 첫째 자리 평균, 정수 최종값)을 돌려준다."""
    if not levels:
        raise ValueError("levels 가 비어 있음")
    mean1 = round_half_up(sum(levels) / len(levels), 1)
    return mean1, int(round_half_up(mean1, 0))


def minutes_between(start, end):
    """시작·종료 시각(time 또는 datetime)의 차이(분). 자정을 넘으면 24시간을 더한다."""
    if start is None or end is None:
        return None
    to_min = lambda t: t.hour * 60 + t.minute + t.second / 60  # noqa: E731
    d = to_min(end) - to_min(start)
    return d + 1440 if d < 0 else d


def diff_stats(diffs):
    """차이 통계: n, 평균, 표본표준편차(n−1), 최대 절댓값의 위치."""
    vals = [d for _, d in diffs]
    n = len(vals)
    if n == 0:
        return {"n": 0, "mean_diff": "", "sd_diff": "", "max_abs_diff": "", "max_abs_point": ""}
    mean = sum(vals) / n
    sd = math.sqrt(sum((v - mean) ** 2 for v in vals) / (n - 1)) if n > 1 else ""
    k = max(range(n), key=lambda i: abs(vals[i]))
    return {"n": n, "mean_diff": round_half_up(mean, 2), "sd_diff": round_half_up(sd, 2) if sd != "" else "",
            "max_abs_diff": round_half_up(abs(vals[k]), 1), "max_abs_point": diffs[k][0]}


# ───────────────────────── 야장 읽기
def _columns(ws):
    return {str(c.value).strip(): c.column for c in ws[HEAD_ROW] if c.value is not None}


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def read_points(wb):
    ws = wb["2_측정점"]
    col = _columns(ws)
    pts = {}
    for r in range(HEAD_ROW + 1, ws.max_row + 1):
        pid = ws.cell(r, col["ID"]).value
        if not pid:
            continue
        pts[str(pid).strip()] = {
            "x": _num(ws.cell(r, col["X(EPSG:5186)"]).value),
            "y": _num(ws.cell(r, col["Y(EPSG:5186)"]).value),
            "z": _num(ws.cell(r, col["높이(m)"]).value),
        }
    return pts


def read_noise(wb):
    """유효 행 목록: dict(point_id, period, Leq, minutes, traffic_ref)."""
    ws = wb["3_소음기록"]
    col = _columns(ws)
    rows = []
    for r in range(HEAD_ROW + 1, ws.max_row + 1):
        g = lambda name: ws.cell(r, col[name]).value  # noqa: E731
        pid, rd, L = g("측정점"), g("회차"), _num(g("LAeq"))
        if not pid or not rd or L is None or str(g("유효(Y/N)") or "").strip().upper() != "Y":
            continue
        start, end = g("시작 시각"), g("종료 시각")
        if isinstance(start, dt.datetime):
            start = start.time()
        if isinstance(end, dt.datetime):
            end = end.time()
        mins = minutes_between(start, end)
        if not mins or mins <= 0:
            continue
        period = "night" if str(rd).strip().upper().startswith("N") else "day"
        rows.append({"point_id": str(pid).strip(), "period": period, "round": str(rd).strip().upper(), "Leq": L,
                     "minutes": mins, "traffic_ref": str(g("교통량 참조") or "").strip()})
    return rows


# ───────────────────────── 집계
OUT_FIELDS = ["point_id", "x", "y", "z", "period", "Leq", "n_samples", "valid_flag", "traffic_count_ref",
              "Leq_legal_mean", "Leq_legal_final", "legal_count_ok", "total_minutes"]


ROUND_FIELDS = ["point_id", "round", "period", "Leq", "minutes", "n_rows", "valid_flag", "traffic_count_ref"]


def combine_rounds(noise):
    """같은 측정점·회차의 행(예: 10분 구간 여러 개)을 측정시간 가중 에너지평균으로 합쳐 회차값을 만든다.
    회차값 = 그 회차 전체 구간의 등가소음도(ES 03300.c 2.10 정의와 같은 계산)."""
    g = {}
    for row in noise:
        g.setdefault((row["point_id"], row["round"]), []).append(row)
    out = []
    for (pid, rd), rows in g.items():
        Ts = [x["minutes"] for x in rows]
        out.append({"point_id": pid, "round": rd, "period": rows[0]["period"],
                    "Leq": energy_average([x["Leq"] for x in rows], Ts), "minutes": sum(Ts), "n_rows": len(rows),
                    "valid_flag": "Y" if sum(Ts) >= MIN_MINUTES else "N",
                    "traffic_count_ref": ";".join(sorted({x["traffic_ref"] for x in rows if x["traffic_ref"]}))})
    return out


def summarize(points, noise):
    """법정 산술평균은 회차값(같은 회차 여러 행은 먼저 합침)으로, 검증용 에너지평균은 모든 행의 측정시간 가중으로."""
    groups, rgroups = {}, {}
    for row in noise:
        groups.setdefault((row["point_id"], row["period"]), []).append(row)
    for rr in combine_rounds(noise):
        rgroups.setdefault((rr["point_id"], rr["period"]), []).append(rr)
    order = list(points) + sorted({p for p, _ in groups} - set(points))
    out = []
    for pid in order:
        for period in ("day", "night"):
            g = groups.get((pid, period))
            if not g:
                continue
            Ls = [x["Leq"] for x in g]
            Ts = [x["minutes"] for x in g]
            rounds = rgroups[(pid, period)]
            mean1, final = arithmetic_average_legal([round_half_up(x["Leq"], 1) for x in rounds])
            p = points.get(pid, {})
            refs = sorted({x["traffic_ref"] for x in g if x["traffic_ref"]})
            out.append({
                "point_id": pid, "x": p.get("x") if p.get("x") is not None else "",
                "y": p.get("y") if p.get("y") is not None else "", "z": p.get("z") if p.get("z") is not None else "",
                "period": period, "Leq": round_half_up(energy_average(Ls, Ts), 1), "n_samples": len(rounds),
                "valid_flag": "Y" if all(x["valid_flag"] == "Y" for x in rounds) else "N",
                "traffic_count_ref": ";".join(refs),
                "Leq_legal_mean": mean1, "Leq_legal_final": final,
                "legal_count_ok": "Y" if len(rounds) >= LEGAL_COUNT[period] else "N",
                "total_minutes": round_half_up(sum(Ts), 1),
            })
    return out


def read_model(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    need = {"point_id", "period", "Leq_calc"}
    if rows and not need <= set(rows[0]):
        raise ValueError(f"모델 결과 CSV에 {sorted(need)} 열이 필요함: {path}")
    out, bad = {}, []
    for line, r in enumerate(rows, start=2):                 # 1행은 머리글
        v = (r["Leq_calc"] or "").strip()
        val = None
        if v and v.lower() != "nan":                          # 빈칸·nan = 계산값 없음(건물 안 등)
            try:
                val = float(v)
            except ValueError:
                bad.append(f"{line}행 point_id={r['point_id']!r} Leq_calc={v!r}")
                continue
            if not math.isfinite(val):
                bad.append(f"{line}행 point_id={r['point_id']!r} Leq_calc={v!r}")
                continue
        rd = (r.get("round") or "").strip().upper()          # round 열이 있으면 회차별 비교
        out[(r["point_id"].strip(), rd) if rd else (r["point_id"].strip(), r["period"].strip())] = val
    if bad:
        raise ValueError(f"모델 결과 CSV의 Leq_calc 가 숫자가 아님({path}):\n  " + "\n  ".join(bad))
    return out


def compare(measured, model, rounds=None):
    """model 키가 (point_id, period)면 시간대별, (point_id, 회차)면 회차별(rounds 필요)로 비교한다."""
    comp, diffs = [], []
    by_round = any(k[1] not in ("day", "night") for k in model)
    items = []
    if by_round:
        for rr in rounds or []:
            items.append(dict(rr, Leq=round_half_up(rr["Leq"], 1), key=(rr["point_id"], rr["round"])))
    else:
        items = [dict(m, key=(m["point_id"], m["period"])) for m in measured]
    for m in items:
        key = m.pop("key")
        calc = model.get(key)
        if calc is None:
            comp.append({"point_id": key[0], "period": m["period"], "round": m.get("round", ""), "Leq_meas": m["Leq"],
                         "Leq_calc": "", "diff": "", "note": "계산값 없음"})
            continue
        d = round_half_up(m["Leq"] - calc, 1)
        comp.append({"point_id": key[0], "period": m["period"], "round": m.get("round", ""), "Leq_meas": m["Leq"],
                     "Leq_calc": calc, "diff": d, "note": "" if m["valid_flag"] == "Y" else "10분 미만 측정 포함"})
        diffs.append((f"{key[0]}/{key[1]}", m["Leq"] - calc))
    return comp, diff_stats(diffs)


def write_csv(path, fields, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def main(argv=None):
    ap = argparse.ArgumentParser(description="야장 xlsx → 측정점별 L_Aeq, 모델 비교")
    ap.add_argument("fieldbook")
    ap.add_argument("--model", help="모델 결과 CSV(point_id, period, Leq_calc[, round]). round 가 있으면 회차별 비교")
    ap.add_argument("--out", default=str(Path(__file__).resolve().parents[1] / "data"))
    a = ap.parse_args(argv)
    wb = load_workbook(a.fieldbook, data_only=False)
    noise = read_noise(wb)
    measured = summarize(read_points(wb), noise)
    rounds = combine_rounds(noise)
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    write_csv(out / "measured_levels.csv", OUT_FIELDS, measured)
    write_csv(out / "measured_rounds.csv", ROUND_FIELDS,
              [dict(r, Leq=round_half_up(r["Leq"], 1), minutes=round_half_up(r["minutes"], 1)) for r in rounds])
    print(f"측정점·시간대 {len(measured)}개 → {out / 'measured_levels.csv'}, 회차 {len(rounds)}개 → measured_rounds.csv")
    if a.model:
        comp, stats = compare(measured, read_model(a.model), rounds)
        write_csv(out / "comparison.csv", ["point_id", "period", "round", "Leq_meas", "Leq_calc", "diff", "note"], comp)
        write_csv(out / "comparison_summary.csv", list(stats), [stats])
        print(f"비교 {stats['n']}개: 평균 차이 {stats['mean_diff']} dB, 표준편차 {stats['sd_diff']} dB, "
              f"최대 {stats['max_abs_diff']} dB ({stats['max_abs_point']})")
    return measured


if __name__ == "__main__":
    main()
